#!/usr/bin/env python3
from __future__ import annotations

import copy
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


SKILL_DIR = Path(
    "/Users/noorjahan/.codex/plugins/cache/openai-primary-runtime/documents/26.430.10722/skills/documents"
)
sys.path.insert(0, str(SKILL_DIR / "scripts"))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402


ACCENT = "1F6B57"
ACCENT_DARK = "184E40"
TEXT = "1F2933"
MUTED = "5B6470"
GRID = "CBD5DF"
HEADER_FILL = "E7F1ED"
CALLOUT_FILL = "F4F8F6"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRID, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def set_paragraph_border(paragraph, color: str = ACCENT, size: str = "10") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), color)


def set_keep_next(paragraph, value: bool = True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if value:
        if keep is None:
            keep = OxmlElement("w:keepNext")
            p_pr.append(keep)
    elif keep is not None:
        p_pr.remove(keep)


def style_font(style, name: str, size_pt: float, color: str = TEXT, bold: bool = False) -> None:
    font = style.font
    font.name = name
    font.size = Pt(size_pt)
    font.color.rgb = RGBColor.from_string(color)
    font.bold = bold
    style.element.rPr.rFonts.set(qn("w:eastAsia"), name)


def tune_styles(doc: Document) -> None:
    styles = doc.styles
    style_font(styles["Normal"], "Arial", 10.2, TEXT)
    normal = styles["Normal"].paragraph_format
    normal.space_before = Pt(0)
    normal.space_after = Pt(5)
    normal.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.line_spacing = 1.08

    style_font(styles["Title"], "Arial", 24, ACCENT_DARK, True)
    title = styles["Title"].paragraph_format
    title.space_before = Pt(0)
    title.space_after = Pt(4)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    style_font(styles["Heading 1"], "Arial", 15.5, ACCENT_DARK, True)
    h1 = styles["Heading 1"].paragraph_format
    h1.space_before = Pt(16)
    h1.space_after = Pt(6)
    h1.keep_with_next = True

    style_font(styles["Heading 2"], "Arial", 12.5, ACCENT, True)
    h2 = styles["Heading 2"].paragraph_format
    h2.space_before = Pt(10)
    h2.space_after = Pt(4)
    h2.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style_font(styles[name], "Arial", 10.1, TEXT)
        fmt = styles[name].paragraph_format
        fmt.left_indent = Inches(0.28)
        fmt.first_line_indent = Inches(-0.18)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(4)
        fmt.line_spacing = 1.05


def text_snapshot(doc: Document):
    paragraphs = [p.text for p in doc.paragraphs]
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
    return paragraphs, tables


def paragraph_has_text(paragraph) -> bool:
    return bool(paragraph.text.strip())


def style_paragraphs(doc: Document) -> None:
    non_empty_seen = 0
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            para.paragraph_format.space_after = Pt(0)
            continue

        non_empty_seen += 1
        style_name = para.style.name if para.style is not None else ""

        for run in para.runs:
            run.font.name = "Arial"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
            if style_name not in ("Title", "Heading 1", "Heading 2"):
                run.font.color.rgb = RGBColor.from_string(TEXT)

        if non_empty_seen == 1:
            para.style = doc.styles["Title"]
            set_paragraph_border(para, ACCENT)
            continue

        if non_empty_seen in (2, 3):
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_after = Pt(7 if non_empty_seen == 2 else 12)
            for run in para.runs:
                run.font.size = Pt(11 if non_empty_seen == 2 else 9.7)
                run.font.color.rgb = RGBColor.from_string(MUTED)
                if non_empty_seen == 2:
                    run.font.bold = True
            continue

        if style_name == "Heading 1":
            set_keep_next(para, True)
        elif style_name == "Heading 2":
            set_keep_next(para, True)
        elif style_name == "Normal":
            para.paragraph_format.space_after = Pt(5)
            para.paragraph_format.line_spacing = 1.08


def column_weights(table) -> list[float]:
    cols = len(table.columns)
    if cols == 1:
        return [1.0]
    weights = []
    for c in range(cols):
        max_len = max((len(row.cells[c].text.strip()) for row in table.rows), default=1)
        avg_len = sum(len(row.cells[c].text.strip()) for row in table.rows) / max(len(table.rows), 1)
        weights.append(max(0.7, min(3.2, 0.55 + (max_len * 0.012) + (avg_len * 0.018))))
    return weights


def style_tables(doc: Document) -> None:
    content_width = (
        doc.sections[0].page_width.twips
        - doc.sections[0].left_margin.twips
        - doc.sections[0].right_margin.twips
    )
    for table in doc.tables:
        rows = len(table.rows)
        cols = len(table.columns)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False

        if cols == 1:
            widths = [content_width]
        else:
            widths = column_widths_from_weights(column_weights(table), total_width_dxa=content_width)

        try:
            apply_table_geometry(
                table,
                widths,
                table_width_dxa=sum(widths),
                indent_dxa=0,
                cell_margins_dxa={"top": 95, "bottom": 95, "start": 130, "end": 130},
            )
        except ValueError:
            pass

        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_border(cell)
                if cols == 1:
                    set_cell_shading(cell, CALLOUT_FILL)
                elif r == 0:
                    set_cell_shading(cell, HEADER_FILL)
                else:
                    set_cell_shading(cell, "FFFFFF")

                for para in cell.paragraphs:
                    para.paragraph_format.space_before = Pt(0)
                    para.paragraph_format.space_after = Pt(3)
                    para.paragraph_format.line_spacing = 1.03
                    if cols > 1 and c == 0 and r > 0:
                        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    elif cols > 1 and len(cell.text.strip()) <= 18:
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in para.runs:
                        run.font.name = "Arial"
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
                        run.font.size = Pt(8.5 if cols >= 4 else 9)
                        run.font.color.rgb = RGBColor.from_string(TEXT)
                        if (cols == 1 and para.text.strip().split("\n", 1)[0]) or (cols > 1 and r == 0):
                            run.font.bold = True
                            run.font.color.rgb = RGBColor.from_string(ACCENT_DARK)

        for row in table.rows:
            row.height = None


def remove_extra_blank_paragraphs(doc: Document) -> None:
    # Preserve all text-bearing paragraphs; only compact empty paragraph spacing.
    for para in doc.paragraphs:
        if not para.text.strip():
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.space_before = Pt(0)


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.start_type = WD_SECTION_START.CONTINUOUS
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.68)
        section.bottom_margin = Inches(0.62)
        section.left_margin = Inches(0.72)
        section.right_margin = Inches(0.72)


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: redesign_railtrack_doc.py input.docx output.docx")

    src = Path(sys.argv[1])
    out = Path(sys.argv[2])
    doc = Document(src)
    before = copy.deepcopy(text_snapshot(doc))

    configure_sections(doc)
    tune_styles(doc)
    style_paragraphs(doc)
    style_tables(doc)
    remove_extra_blank_paragraphs(doc)

    after = text_snapshot(doc)
    if before != after:
        raise RuntimeError("Text changed during redesign; refusing to write output.")

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
